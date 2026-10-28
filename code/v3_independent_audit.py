#!/usr/bin/env python3
"""Read-only, independent numerical and chronological checks for revision v3.

This script does not import the fitting/feature-construction implementation.
Optional --replay compares newly executed stage CSVs with preserved outputs;
optional --manuscript compares all displayed table cells with paper_tables.
It does not authenticate historical source vintages or establish field efficacy.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, math
from pathlib import Path
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "support/baseline/base_v11"
REGIONS = ("KR011","KR012","KR013","KR021","KR022","KR023","KR031","KR032",
           "KR041","KR042","KR043","KR051","KR052","KR053","KR054","KR061","KR071")
BOUNDARIES = {"KR053": {2017}, "KR031": {2024}, "KR032": {2024}}

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out",type=Path,required=True)
    ap.add_argument("--replay",type=Path)
    ap.add_argument("--manuscript",type=Path)
    args=ap.parse_args()
    checks=[]
    def check(name,condition,**evidence):
        checks.append(dict(check=name,passed=bool(condition),**evidence))
    def close(a,b):return bool(np.allclose(a,b,rtol=0,atol=1e-10,equal_nan=True))
    source=pd.read_csv(BASE/'source_harmonized/kosis_ksic11_H1_2013_2025_panel.csv',float_precision='round_trip')
    panel=pd.read_csv(BASE/'analysis_harmonized/results/main/analysis_panel.csv',float_precision='round_trip')
    lookup=source.set_index(['region_code','year','industry_code']).employment_persons.to_dict()
    def E(r,y,i):return lookup.get((r,y,i),float('nan'))
    def good(v):return bool(np.isfinite(v) and v>0)
    errors={k:[] for k in ['target_logpoints','log_employment','industry_share_pct',
        'own_growth','own_growth_lag1','peer_industry_growth','region_total_growth',
        'lagged_available_region_industry_growth']}
    boundary_valid=True;peer_keys=True;years=True
    for row in panel.itertuples():
        r,y,i=row.region_code,int(row.target_year),row.industry_code
        past=E(r,y-1,i);older=E(r,y-2,i);oldest=E(r,y-3,i)
        peer=[rr for rr in REGIONS if rr!=r and y-1 not in BOUNDARIES.get(rr,set())
              and good(E(rr,y-1,i)) and good(E(rr,y-2,i))]
        all_valid=[rr for rr in REGIONS if y-1 not in BOUNDARIES.get(rr,set())
              and good(E(rr,y-1,i)) and good(E(rr,y-2,i))]
        values={
          'target_logpoints':100*math.log(E(r,y,i)/past),
          'log_employment':math.log(past),
          'industry_share_pct':100*past/E(r,y-1,'_T'),
          'own_growth':100*math.log(past/older),
          'own_growth_lag1':100*math.log(older/oldest),
          'peer_industry_growth':100*math.log(sum(E(rr,y-1,i) for rr in peer)/sum(E(rr,y-2,i) for rr in peer)),
          'region_total_growth':100*math.log(E(r,y-1,'_T')/E(r,y-2,'_T')),
          'lagged_available_region_industry_growth':100*math.log(sum(E(rr,y-1,i) for rr in all_valid)/sum(E(rr,y-2,i) for rr in all_valid))}
        for k,v in values.items():errors[k].append(abs(v-getattr(row,k)))
        boundary_valid &= not any(t in BOUNDARIES.get(r,set()) for t in (y,y-1,y-2))
        peer_keys &= row.peer_region_codes==';'.join(peer)
        years &= row.input_year==y-1
    for k,v in errors.items():check('independent raw feature '+k,max(v)<1e-10,rows=len(v),max_abs_difference=max(v))
    check('all retained rows satisfy stated boundary exclusion',boundary_valid)
    check('peer region identifiers match independent eligible sets',peer_keys)
    check('input year is target year minus one',years)
    check('chronological split sizes',panel.groupby('split').size().to_dict()=={'train':558,'validation':463,'test':428})
    check('sample inclusion 1449 of 1700',len(panel)==1449,inclusion_pct=100*len(panel)/1700)
    pred=pd.read_csv(BASE/'analysis_harmonized/results/main/test_predictions.csv',float_precision='round_trip')
    ridge=np.abs(pred.target_logpoints-pred.ridge_alpha1000)
    ref=np.abs(pred.target_logpoints-pred.global_past_mean)
    delta=float((ridge-ref).mean())
    check('initial MAE difference rounds to -0.106',round(delta,3)==-.106,actual=delta)
    check('428 predictions are the test panel',len(pred)==428 and pred.target_year.min()==2023 and pred.target_year.max()==2025)
    check('all initial global mean forecasts use development targets only',close(pred.global_past_mean,panel[panel.split!='test'].target_logpoints.mean()))
    bs=pd.read_csv(ROOT/'results/forecast_replay/all_paired_comparisons.csv')
    row=bs[(bs.subset=='all') & (bs.model=='ridge_alpha1000')].iloc[0]
    check('initial interval includes zero',row.ci95_lo<0<row.ci95_hi,ci95=[float(row.ci95_lo),float(row.ci95_hi)])
    exp=pd.read_csv(ROOT/'results/benchmark/expanding_comparisons.csv')
    for name,should_exclude in [('learned32',True),('all38',False)]:
        z=exp[(exp.model==name)&(exp.metric=='mae')].iloc[0]
        excludes=bool(z.sim_hi<0 or z.sim_lo>0)
        check('expanding simultaneous interval '+name,excludes==should_exclude,
              delta=float(z.delta),interval=[float(z.sim_lo),float(z.sim_hi)])
    field=pd.read_csv(ROOT/'field_study/responses_EMPTY.csv')
    schedule=pd.read_csv(ROOT/'field_study/counterbalanced_assignment.csv')
    check('field response file has no observations',len(field)==0)
    check('field assignment is a plan of 8 participants x 16 tasks',len(schedule)==128 and schedule.participant_code.nunique()==8 and schedule.task_id.nunique()==16)
    transfer=pd.read_csv(ROOT/'results/transfer_evaluation/case_results.csv.gz')
    check('transfer judgments 30360',len(transfer)==30360)
    pivot=transfer.pivot(index='case_id',columns='engine',values='permitted')
    check('7590 matched native and Pandera decisions',len(pivot)==7590 and pivot.native_registered.equals(pivot.pandera_registered))
    trust=pd.read_csv(ROOT/'results/transfer_evaluation/trust_boundary_results.csv')
    # Trust probes may be stored wide or long; expected release is wide.
    cols=[c for c in trust if c.endswith('_permitted')]
    check('trust-boundary accepted probes retained',len(trust)==506 and bool(trust[cols].all().all()) if cols else False,rows=len(trust),decision_columns=cols)
    table_cells=0
    table_changes=[]
    presentation_map={'\u2212':'-','\u2013':'-','\u2014':'-'}
    if args.manuscript:
        from docx import Document
        doc=Document(args.manuscript)
        index=json.loads((ROOT/'results/paper_tables/index.json').read_text())
        tables=[x for x in index['records'] if x['type']=='table']
        offset=1 if len(doc.tables)==len(tables)+1 else 0
        check('24 research tables, with an optional abstract layout table',len(doc.tables)-offset==len(tables)==24)
        for dtable,record in zip(doc.tables[offset:],tables):
            expected=list(csv.reader((ROOT/record['path']).open(encoding='utf-8-sig',newline='')))
            actual=[[c.text for c in row.cells] for row in dtable.rows]
            table_cells+=sum(len(row) for row in actual)
            raw_equal=actual==expected
            normalize=lambda rows:[[''.join(cell.translate(str.maketrans(presentation_map)).split()) for cell in row] for row in rows]
            normalized_equal=normalize(actual)==normalize(expected)
            changed=[]
            for ri,(a_row,e_row) in enumerate(zip(actual,expected)):
                for ci,(a_cell,e_cell) in enumerate(zip(a_row,e_row)):
                    if a_cell!=e_cell:
                        item={'table_id':record['id'],'row':ri,'column':ci,
                              'csv_text':e_cell,'manuscript_text':a_cell,
                              'presentation_only':''.join(a_cell.translate(str.maketrans(presentation_map)).split())==''.join(e_cell.translate(str.maketrans(presentation_map)).split())}
                        changed.append(item)
                        table_changes.append(item)
            check('displayed table '+record['id'],normalized_equal,raw_equal=raw_equal,
                  changed_cells=len(changed),allowed_normalization='U+2212/U+2013/U+2014 to ASCII hyphen; layout whitespace only')
    replay=[]
    if args.replay:
        status=json.loads((args.replay/'run_summary.json').read_text())
        check('nine-stage replay completed',status['status']=='passed' and len(status['jobs'])==9)
        for p in sorted(args.replay.rglob('*')):
            if not p.is_file() or not (p.name.endswith('.csv') or p.name.endswith('.csv.gz')):continue
            rel=p.relative_to(args.replay)
            q=ROOT/'results'/rel
            if not q.exists():
                gz=q.with_name(q.name+'.gz')
                if gz.exists():q=gz
                else:continue
            try:
                a=pd.read_csv(p,float_precision='round_trip');b=pd.read_csv(q,float_precision='round_trip')
                pd.testing.assert_frame_equal(a,b,check_dtype=False,check_exact=False,rtol=1e-10,atol=1e-10)
                replay.append({'path':str(rel),'passed':True,'rows':len(a)})
            except Exception as exc:replay.append({'path':str(rel),'passed':False,'error':str(exc)[:1500]})
        check('all corresponding rerun CSVs match preserved results',bool(replay) and all(x['passed'] for x in replay),files=len(replay))
    report={'status':'PASS' if all(c['passed'] for c in checks) else 'FAIL',
       'checks_passed':sum(c['passed'] for c in checks),'checks_total':len(checks),
       'checks':checks,'displayed_cells':table_cells,'displayed_cell_changes':table_changes,
       'presentation_character_map':presentation_map,'replay_csv_comparison':replay,
       'scope':'Numerical consistency, source-derived feature timing, preserved/replayed results; not source authenticity or human efficacy.'}
    report['audit_script_sha256']=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    if args.manuscript:
        report['manuscript']={'file_name':args.manuscript.name,
            'bytes':args.manuscript.stat().st_size,
            'sha256':hashlib.sha256(args.manuscript.read_bytes()).hexdigest()}
    args.out.parent.mkdir(parents=True,exist_ok=True)
    args.out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'status':report['status'],'checks':len(checks),'failed':[c for c in checks if not c['passed']]},ensure_ascii=False,indent=2))
    if report['status']!='PASS':raise SystemExit(1)

if __name__=='__main__':main()
