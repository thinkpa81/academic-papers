from pathlib import Path
import json,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
import argparse,hashlib
parser=argparse.ArgumentParser(description='Rebuild the verified 2026-10-28 manuscript from the preserved source template.')
parser.add_argument('--output',type=Path,default=ROOT/'generated/00_final_manuscript_261028.docx')
args=parser.parse_args()
source=ROOT/'support/manuscript/source_before_revision_20261028.docx'
assert hashlib.sha256(source.read_bytes()).hexdigest()=='c9783542ed8520034b52b5f0a572c49359029311cf6a82cae0d728a12930d91e', 'Unexpected source manuscript'
d=Document(source)
changes=[]
def setp(p,text,location):
 old=p.text
 rpr=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
 changes.append({'location':location,'before':old,'after':text})
def replace(i,text):setp(d.paragraphs[i],text,f'paragraph_{i}')
replace(5,'Purpose  This study designs a reviewable regional employment decision support service that links data eligibility, comparative forecasts, and current review records at the prediction level.')
replace(6,'Design/methodology/approach  The design specifies four principles and an evidence-linked review workflow. Retrospective forecasting used harmonized first-half statistics for 17 Korean provinces and 10 industry groups, with 428 evaluation observations for 2023–2025. Fault injection, ablation, model substitution, and matched-information comparisons with Pandera evaluated the specified review requirements; selected requirements were also applied to three US macroeconomic series.')
replace(7,'Findings  The initial ridge–benchmark MAE difference was −0.106 log points, with a province-cluster interval including zero; expanding-window results depended on the selection strategy. Linking current evidence to review records increased blocked faults from 16 to 36 out of 36 specified cases while accepting all 14 controls. Registered-reference checks blocked 14 known inconsistencies and accepted 272 consistent inputs. Native and Pandera implementations agreed on 7,590 cases under matched requirements and information. Joint alteration of submitted evidence and trusted references passed in all 506 separate probes. The contribution is an executable specification of prediction-level review conditions and its tested operating boundaries. Practitioner judgment and organizational outcomes require subsequent field evaluation.')
replace(8,'Keywords: Regional employment, Decision support service, Design science, Reviewability, Data eligibility, Evidence-linked review')
summary=d.tables[0].cell(0,0).paragraphs[0]
new_summary='본 연구는 지역고용 예측의 자료 사용조건, 모형 비교근거와 현재 검토기록을 예측 단위에서 연결하는 의사결정지원서비스를 설계·평가하였다. 검토가능성을 자료 적격성 확인부터 근거 제시, 검토 완료 및 근거 변경 시 재검토까지의 절차로 구체화하였다. 17개 시도·10개 산업의 상반기 자료에서 2023~2025년 428개 관측치를 평가한 결과, 최초 릿지와 전체 과거평균의 MAE 차이는 −0.106로그포인트였으며 시도 군집 구간은 0을 포함하였다. 관계 검사를 추가하면 명세한 결함 36개 중 차단 수가 16개에서 36개로 늘었고 정상 14개는 모두 수용하였다. 동일 요건·정보의 Pandera 비교와 미국 공개계열 적용에서는 공통 검사 요건의 실행 가능성을 확인하였으며, 참조자료까지 함께 바꾸면 검사를 통과하는 경계도 확인하였다. 공헌은 지역고용 예측의 검토 완료 요건과 적용조건을 실행 가능한 명세로 제시한 데 있다. 실제 담당자의 판단 개선과 업무성과는 후속 현장평가의 대상이다.'
old_summary=summary.text
summary.runs[1].text='\n'+new_summary
summary.runs[2].text=''
summary.runs[3].text=''
summary.runs[4].text='\nKeywords: 지역고용, 의사결정지원서비스, 설계과학, 검토가능성, 자료 적격성, 예측근거'
changes.append({'location':'Korean abstract','before':old_summary,'after':summary.text})
replace(16,'본 연구의 핵심 기여는 예측별 자료 사용조건·비교근거·현재 검토기록을 결합하고, 그 관계가 충족될 때 연구용 검토를 완료하도록 명세한 서비스 설계에 있다. 여기서 검토가능성은 담당자가 개별 예측의 근거와 적용조건을 확인하고, 판단사유를 현재 근거에 연결하며, 근거 변경 후 재검토 대상을 식별할 수 있도록 하는 설계 특성을 뜻한다[55]. 이 기여를 뒷받침하기 위해 시간순 예측평가로 적용조건을 도출하고, 관계 검사와 등록 근거 대조를 구현하였으며, 기능 제거·모형 교체·같은 정보의 도구 비교와 다른 집계자료 적용으로 작동범위를 확인하였다. 적용 맥락은 자료판본·예측기간과 참조자료의 갱신권한을 구분해 관리할 수 있는 집계통계 검토업무이다.')
replace(32,'본 연구는 업무 문제를 해결하는 산출물을 설계·평가하는 설계과학 연구(design science research, DSR)의 관점을 적용하였다[31,44]. 산출물은 예측근거 카드, 권역–산업 비교표 및 현재 근거에 연결된 검토기록이고, 이를 묶는 검토 완료 요건을 설계 지식의 단위로 설정하였다[53]. 설계원칙(design principle, DP)의 목표·맥락·작동방식·근거를 명세하는 접근[30]에 따라 업무 과업, 구현 기능과 평가항목을 대응시켰다. 인공적·형성적 평가는 명세한 관계의 작동을, 후속 현장평가는 담당자 판단과 업무 유용성을 검토하도록 구분하였다[49].')
replace(40,'<그림 1>은 담당자의 과업과 그 과업을 뒷받침하는 근거 연결을 함께 나타낸다. DP2는 자료 사용조건을, DP1·DP3은 비교예측과 적용범위를, DP4는 현재 근거에 대한 검토기록을 다룬다. 별도로 등록한 모형·입력·검증요약과의 대조를 거쳐 연구용 검토 상태를 기록하고, 근거가 바뀌면 같은 예측 식별자에 새 판본을 연결하여 재검토한다. 예측평가와 기술시험은 이 흐름의 적용조건 및 작동근거를 제공한다. RSC 수요조사와 기관 승인은 현장 적용 단계이며 담당자 효과는 후속 평가로 구분하였다.')
replace(42,'<그림 1> 검토가능성을 위한 지역고용 의사결정지원서비스의 과업과 근거 연결')
replace(43,'주: 실선은 구현·기술평가 범위, 점선은 후속 현장 업무 연계를 나타낸다. 등록 근거의 신뢰성과 갱신권한은 기관 운영조건이며, 연구용 검토 완료와 정책 승인은 구분된다.')
# Clarify which materials are published and which accompany this revision.
idx=next(i for i,p in enumerate(d.paragraphs) if p.text.startswith('분석 코드, 원자료 확보본,'))
replace(idx,'기존 예측분석과 서비스 기술평가의 코드·원자료 확보본·계산 결과는 GitHub 저장소(https://github.com/thinkpa81/academic-papers)에 제공하였다. 이번 원고와 함께 제공하는 재현패키지에는 원자료 정규화, 예측 재적합, 후보군·재선정 분석, 등록 근거 및 공통 요건 비교·이식시험의 실행 기록과 파일 해시를 수록하였다. 2026년 10월 28일 보완한 그림 1의 생성 코드와 현장평가용 과업·수집 프로그램·분석 코드는 동봉 재현패키지에 포함하였다. 현장평가 자료는 실시 계획과 빈 관측양식이며, 실제 결과와 구분하였다. 원자료와 시험의 상세 명세는 부록 B, E와 G에 제시하였다.')
replace(230,'5.2 검토가능성의 명세와 서비스 설계기여')
replace(231,'서비스 설계의 공헌은 <그림 1>의 과업 흐름에서 개별 예측의 검토 완료 요건을 구체화한 데 있다. 담당자가 확인할 비교기준·검증사례 수·집계범위를 카드에 묶고, 판단사유가 현재 자료판본과 등록 모형을 참조하는지 검사하도록 구성하였다. 그 결과 동일한 예측값을 표시하더라도 근거 판본이나 검토기록의 연결이 달라지면 서로 다른 검토 상태를 부여할 수 있다. 이는 문서화[20,29,39], 검토 가능성[45,55], 출처 일관성과 자료 검증[56–58]의 원리를 지역고용 업무의 과업·입력·판정·기록 관계로 구체화한 적용 지식이다.')
replace(232,'이 설계의 기술적 작동근거는 동일 입력 비교와 기능 제거에서 확인하였다. 관계 검사를 추가하면 초기 결함 36개 중 차단 수가 16개에서 36개로 늘었고, 시점·수치·연결 검사를 각각 제거하면 12·8·9개 결함이 통과하였다(<표 12>). 모형정보를 함께 바꾼 14개는 등록 근거 대조에서 차단되었다(<표 14>). 같은 요건·정보를 제공한 Pandera도 직접 구현과 같은 판정을 산출하였다(<표 15>). 따라서 재사용할 설계 지식은 업무별 관계 요건과 필요한 참조정보의 명세이며, 검증 라이브러리는 이를 실행하는 수단으로 해석한다.')
replace(236,'서울 BTE 사례의 사용 시나리오에서는 릿지 −0.40, 기준모형 +0.67 및 편차 −1.07로그포인트와 유효 검증사례 0개를 함께 확인한다(<그림 2>). 담당자의 후속 과업은 구성자료와 산업분류의 연속성을 점검하고, 미충원 직종·채용계획·직무전환 수요를 현장조사 가설로 작성하는 것이다. 기록에는 확인할 자료, 판단사유, 담당자와 현재 근거판본을 남긴다. 자료가 보완되면 예측근거를 갱신하고 이전 기록의 재검토 여부를 확인한다. 이 절차는 시제품의 업무 연계 시나리오이며, 실제 담당자가 수행한 행동이나 훈련수요의 정답을 관측한 결과와 구분된다.')
replace(244,'본 연구는 자료 적격성 확인, 예측 비교, 근거 제시와 현재 검토기록을 하나의 과업 흐름으로 연결하여 지역고용 예측의 검토가능성을 구체화하였다. RQ1에서 최초 릿지의 MAE 차이는 −0.106로그포인트였고 신뢰구간은 0을 포함하였다. 확장창의 학습후보 32개 선택에서는 기준 6개 선택 대비 차이 −0.270로그포인트의 동시구간이 [−0.503, −0.037]이었으나, 전체 38개 선택에서는 0을 포함하였다. 이 탐색적 결과는 성능을 후보군·선정규칙·평가지표 및 적용기간과 함께 제시할 필요성을 뒷받침한다.')
replace(245,'RQ2에서는 현재 근거에 연결된 기록과 등록 참조자료의 일치를 연구용 검토 완료 요건으로 명세하였다. RQ3에서는 관계 검사·기능 제거·모형 교체와 도구·자료 간 적용으로 그 작동범위를 확인하였다. 동일 요건·정보의 구현 간 판정 일치와 참조자료 동시 변경 시의 통과는 검사 요건 및 신뢰조건의 역할을 함께 보여준다. 본 연구는 지역고용 검토업무의 실행 가능한 관계 명세와 검토 상태 관리 절차를 제시하였으며, 담당자의 판단과 업무성과는 후속 현장평가에서 검증할 과제로 구분하였다.')
replace(252,'사용자 평가용으로 실제 예측 수치와 실험용 검토 상태를 결합한 16개 과업 및 실무자 8명 규모의 예비평가 계획을 준비하였다. 표·카드의 두 조건에 동일한 정보항목을 제시하고, 참여자별 과업 유형과 조건을 균형 배정하였다. 판단 정확도와 최초 판단시간을 일차·이차 지표로 기록하며 근거 신뢰성은 유형별 탐색 문항으로 측정한다. 수집 프로그램과 참여자별 비교 코드는 부록 F의 절차로 실행할 수 있다. 실제 참여자와 응답은 0이며, 확증평가의 표본수는 예비평가의 변동과 최소 실무적 차이에 근거해 정하도록 하였다.')
replace(378,'현장평가 준비자료는 실제 예측 수치에 실험용 검토 상태를 결합한 16개 과업, 표·카드의 두 표시조건, 계획 인원 8명의 배정표 및 빈 관측·채점양식으로 구성하였다. 각 참여자는 조건별 8개 과업을 수행하며 정상, 이전 판본, 검증사례 수 불일치와 모형계수 불일치의 네 유형이 조건별로 2개씩 포함된다. 과업별로 두 조건에 각각 4회 배정하여 같은 과업의 재노출을 피하고 정보항목을 동일하게 제시하도록 준비하였다. 예측값·기준값·원사례 식별자와 실험용 변경내용은 함께 보존하였다.')
replace(379,'로컬 수집 프로그램은 과업 시작부터 최초 판단 제출까지의 시간을 기록하고, 이후 판단사유·확신도·근거 신뢰성을 별도로 받는다. 정답을 알려 주는 결함명과 피드백은 최초 판단 전에 제시하지 않는다. 분석 코드는 참여자별 정답률·시간 중앙값과 카드−표의 차이를 산출하며 반복행을 독립 참여자로 계산하지 않는다. 1~7점 근거 신뢰성은 탐색적 단일 문항이다. 이 준비자료는 표시 구성의 예비평가를 위한 것이며 전체 서비스의 업무효과를 검증한 결과는 아니다. 실제 참여자와 관측행은 0이고 빈 양식의 분석 상태는 not_evaluated이다. 실행 절차와 과업·수집·분석 코드는 재현패키지 field_study 및 code에 수록하였다.')
# Replace only the first embedded figure, preserving all equations and other images.
shape=d.inline_shapes[0];rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
picture=ROOT/'figures/paper_figure_01.png'
d.part.related_parts[rid]._blob=picture.read_bytes()
w,h=Image.open(picture).size;shape.height=int(shape.width*h/w)
for highlight in list(d.element.iter(qn('w:highlight'))):
    highlight.getparent().remove(highlight)
for row in d.tables[10].rows:
    for cell in row.cells:
        for side in ('top','bottom'):
            cell._tc.tcPr.find(qn('w:tcMar')).find(qn('w:'+side)).set(qn('w:w'),'20')
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(0.5)
            p.paragraph_format.space_after=Pt(0.5)
out=args.output;out.parent.mkdir(parents=True,exist_ok=True);d.save(out)
out.with_suffix('.changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print('changed paragraphs',len(changes),'tables',len(d.tables),'figures',len(d.inline_shapes))
