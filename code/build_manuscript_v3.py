#!/usr/bin/env python3
"""Build the anonymous 261028_v3 DOCX from the preserved source and audit spec.

The supplied HWP template determines geometry and core typography. The available
NanumMyeongjo replaces proprietary HWP fonts. This writes genuine DOCX, not HWP.
Native Word equations, scientific tables and all numeric results are preserved.
"""
from pathlib import Path
import argparse,copy,hashlib,io,json,re,zipfile
from lxml import etree
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DASH_MAP=str.maketrans({c:'-' for c in '\u2212\u2010\u2011\u2012\u2013\u2014\ufe63\uff0d'})

def plain(p):return ''.join(p._p.xpath('.//w:t/text()'))

def replace_text(p,old,new):
    """Replace visible text without dropping native equations or superscripts."""
    nodes=p._p.xpath('.//w:t')
    joined=''.join(n.text or '' for n in nodes)
    if old==new:return
    assert old in joined,(old,joined)
    start=joined.index(old);end=start+len(old);pos=0;first=True
    for n in nodes:
        value=n.text or '';nend=pos+len(value)
        if nend>start and pos<end:
            left=max(0,start-pos);right=min(len(value),end-pos)
            n.text=value[:left]+(new if first else '')+value[right:]
            n.set('{http://www.w3.org/XML/1998/namespace}space','preserve');first=False
        pos=nend
    assert ''.join(n.text or ''for n in nodes)==joined[:start]+new+joined[end:]

def set_simple(p,text):
    assert not p._p.xpath('.//m:oMath'),'Do not flatten native math'
    assert not any(r.font.superscript for r in p.runs),'Preserve superscript via localized replacements'
    rp=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(text)
    if rp is not None:r._r.insert(0,rp)

def prop(rp,tag,val):
    e=rp.find(qn('w:'+tag))
    if e is None:e=OxmlElement('w:'+tag);rp.append(e)
    e.set(qn('w:val'),str(val))

def apply_typography(doc):
    for style in doc.styles:
        if style.type==1:
            style.font.name='NanumMyeongjo'
            rp=style.element.get_or_add_rPr()
            rf=rp.find(qn('w:rFonts'))
            if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
            for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'NanumMyeongjo')
            prop(rp,'w',95)
    for section in doc.sections:
        section.page_width=Mm(190);section.page_height=Mm(262)
        section.left_margin=Mm(20);section.gutter=Mm(3);section.right_margin=Mm(20)
        section.top_margin=Mm(38);section.bottom_margin=Mm(22)
        section.header_distance=Mm(20);section.footer_distance=Mm(10)
        cols=section._sectPr.find(qn('w:cols'))
        if cols is not None:cols.set(qn('w:space'),str(round(10/25.4*1440)))
    # Retain source geometry for special headings/notes; main type follows template.
    doc.styles['Normal'].font.size=Pt(9.5)
    for i,p in enumerate(doc.paragraphs):
        text=p.text
        if i>=10 and p.style.name=='Normal' and text and not text.startswith(('<','주:','[')):
            if not p._p.xpath('.//w:drawing') and not p._p.xpath('.//m:oMath'):
                p.paragraph_format.line_spacing=Pt(15.2)
                p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
        if p.style.name=='Heading 1':
            p.paragraph_format.line_spacing=Pt(15.6)
            for r in p.runs:r.font.size=Pt(13)
            for old,new in [('I.','1.'),('II.','2.'),('III.','3.'),('IV.','4.'),('V.','5.'),('VI.','6.')]:
                if p.text.startswith(old+' '):replace_text(p,old+' ',new+' ');break
        if re.match(r'^\[\d+\]',text):
            p.paragraph_format.line_spacing=Pt(15.2)
            p.paragraph_format.left_indent=Mm(7.0556)
            p.paragraph_format.first_line_indent=Mm(-7.0556)
            p.paragraph_format.space_after=Pt(2)
    title=doc.paragraphs[0]
    title.paragraph_format.line_spacing=Pt(27.2)
    title.paragraph_format.space_before=Pt(12)
    title.paragraph_format.space_after=Pt(20)
    for r in title.runs:r.font.size=Pt(17)
    en=doc.paragraphs[1]
    en.paragraph_format.line_spacing=Pt(18)
    en.paragraph_format.space_after=Pt(20)
    for r in en.runs:r.font.size=Pt(12)
    # The clean source uses one cell for heading, summary, and keywords. Separate
    # these into paragraphs inside that same cell to match the supplied box.
    c=doc.tables[0].cell(0,0)
    content=c.text.split('\n')
    assert len(content)==3 and content[0]=='요 약'
    c.text=''
    p=c.paragraphs[0];p.add_run(content[0]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=Pt(12);p.paragraph_format.space_after=Pt(8)
    for r in p.runs:r.font.size=Pt(9);r.bold=True
    p=c.add_paragraph(content[1]);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=Pt(18);p.paragraph_format.first_line_indent=Mm(3.5278)
    p.paragraph_format.left_indent=Mm(4.4169);p.paragraph_format.right_indent=Mm(4.4169)
    for r in p.runs:r.font.size=Pt(10)
    p=c.add_paragraph(content[2]);p.paragraph_format.space_before=Pt(10)
    p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.left_indent=Mm(4.4169)
    p.paragraph_format.right_indent=Mm(4.4169)
    for r in p.runs:r.font.size=Pt(10)
    # Keep the multi-reference label together instead of splitting the digits 39.
    label=doc.tables[1].cell(2,0).paragraphs[0]
    label_text=label.text.replace('FactSheets[20,29,39]','FactSheets\n[20,29,39]')
    if label_text!=label.text:
        label.clear();label.add_run(label_text)
    # Scale existing full-width tables to the template's 147 mm body width.
    for table in doc.tables:
        for p in table._tbl.xpath('.//w:p'):
            pp=p.find(qn('w:pPr'))
            if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
            wrap=pp.find(qn('w:wordWrap'))
            if wrap is None:wrap=OxmlElement('w:wordWrap');pp.append(wrap)
            wrap.set(qn('w:val'),'0')
        grid=table._tbl.tblGrid.gridCol_lst
        total=sum(int(e.get(qn('w:w'))) for e in grid)
        maximum=round(147/25.4*1440)
        if total>maximum:
            factor=maximum/total
            for e in grid:e.set(qn('w:w'),str(round(int(e.get(qn('w:w')))*factor)))
            for e in table._tbl.xpath('.//w:tcW'):
                if e.get(qn('w:type'))=='dxa':e.set(qn('w:w'),str(round(int(e.get(qn('w:w')))*factor)))
        # No exact row heights: keep content readable after width adjustment.
        for e in table._tbl.xpath('.//w:trHeight'):
            e.set(qn('w:hRule'),'atLeast')
    for p in doc.paragraphs:
        if not p.text and not p._p.xpath('.//w:drawing') and not p._p.xpath('.//m:oMath'):
            if p is doc.paragraphs[4]:p.paragraph_format.space_after=Pt(0)
    # Remove the previous running title; no author/affiliation header is added.
    seen=set()
    for section in doc.sections:
        for header in [section.header,section.first_page_header,section.even_page_header]:
            if id(header._element) in seen:continue
            seen.add(id(header._element))
            for p in header.paragraphs:
                if 'Regional Employment Decision Support' in p.text:
                    replace_text(p,'Regional Employment Decision Support','')
            for node in header._element.xpath('.//w:t'):
                if node.text:node.text=node.text.replace('Regional Employment Decision Support','')

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output',type=Path,default=ROOT/'generated/00_최종논문_261028_v3.docx')
    ap.add_argument('--code-commit',required=True,help='Verified immutable code commit SHA; draft builds may use 40 zeroes')
    args=ap.parse_args();assert re.fullmatch(r'[0-9a-f]{40}',args.code_commit)
    spec=json.loads((ROOT/'support/manuscript/revision_261028_v3.json').read_text())
    source=ROOT/spec['source'];assert hashlib.sha256(source.read_bytes()).hexdigest()==spec['source_sha256']
    doc=Document(source)
    maths=[etree.tostring(e) for e in doc._element.xpath('.//m:oMath')]
    for edit in spec['paragraphs']:
        p=doc.paragraphs[edit['index']];assert p.text==edit['before'],edit['index']
        after=edit['after'].replace('{CODE_COMMIT}',args.code_commit)
        if edit.get('replacements') and after==edit['after']:
            for change in edit['replacements']:replace_text(p,change['old'],change['new'])
        else:
            if any(r.font.superscript for r in p.runs):
                # These edits keep the tail containing the native exponent.
                old,new=edit['before'],after
                from difflib import SequenceMatcher
                for tag,a,b,c,d in reversed(SequenceMatcher(None,old,new,autojunk=False).get_opcodes()):
                    if tag=='equal':continue
                    if a==b:
                        anchor=old[max(0,a-1):a];replace_text(p,anchor,anchor+new[c:d])
                    else:replace_text(p,old[a:b],new[c:d])
            else:set_simple(p,after)
        assert p.text==after,(edit['index'],p.text,after)
    for edit in spec['cells']:
        c=doc.tables[edit['table']].cell(edit['row'],edit['column'])
        assert c.text==edit['before'];assert len(c.paragraphs)==1
        set_simple(c.paragraphs[0],edit['after'])
    for edit in spec['figures']:
        shape=doc.inline_shapes[edit['index']]
        rid=shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
        doc.part.related_parts[rid]._blob=(ROOT/edit['path']).read_bytes()
        factor=min(1,Mm(147)/shape.width)
        shape.width=int(shape.width*factor);shape.height=int(shape.height*factor)
        for e in shape._inline.xpath('.//a:xfrm/a:ext'):
            e.set('cx',str(shape.width));e.set('cy',str(shape.height))
    apply_typography(doc)
    count=0
    for e in doc._element.xpath('.//w:t'):
        old=e.text or '';new=old.translate(DASH_MAP);count+=sum(a!=b for a,b in zip(old,new));e.text=new
    assert maths==[etree.tostring(e) for e in doc._element.xpath('.//m:oMath')],'Native equations changed'
    assert not any(chr(c) in ''.join(doc._element.xpath('.//w:t/text()')) for c in DASH_MAP)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    stream=io.BytesIO()
    with zipfile.ZipFile(args.output) as src,zipfile.ZipFile(stream,'w',compression=zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info=zipfile.ZipInfo(name,(2026,10,28,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;dst.writestr(info,src.read(name))
    args.output.write_bytes(stream.getvalue())
    audit={'release':'261028_v3','code_commit':args.code_commit,'source_sha256':spec['source_sha256'],
           'paragraph_edits':len(spec['paragraphs']),'cell_edits':len(spec['cells']),'native_equations_preserved':len(maths),
           'tables':len(doc.tables),'figures':len(doc.inline_shapes),'ascii_dash_replacements':count,
           'font_substitution':'NanumMyeongjo','format':'DOCX, supplied HWP template geometry',
           'sha256':hashlib.sha256(args.output.read_bytes()).hexdigest()}
    args.output.with_suffix('.build.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(audit,ensure_ascii=False))

if __name__=='__main__':main()
